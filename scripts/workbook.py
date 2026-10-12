#!/usr/bin/env python3
"""Compile numerical supplementary tables and their definitions into a workbook."""
from __future__ import annotations

import csv
import hashlib
import json
import math
import re
import shutil
import textwrap
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.workbook.properties import CalcProperties

ROOT = Path(__file__).resolve().parents[1]
REV = ROOT
NAME = "supplementary_tables.xlsx"
NAVY = "233647"
GRAY = "F3F5F7"
INK = "17212B"
MUTED = "53616F"
LINE = Side(style="thin", color="D9E0E6")
FONT = "Arial"
SCOPES = {
    1: "Observed coverage before filling or forecast evaluation.",
    2: "Predictors and indicator combinations used in each model.",
    3: "Common forecast weeks; full training histories; final four outcome observations excluded.",
    4: "Conditional comparisons on common forecast weeks; full training histories; final four outcome observations excluded.",
    5: "One-week comparisons against AR; full training histories; final four outcome observations excluded.",
    6: "One-week comparisons; full training histories; final four outcome observations retained.",
    7: "One-week calibration on common forecast weeks; final four outcome observations excluded.",
    8: "Identical training observations within each contrast; final four outcome observations excluded.",
    9: "One-week comparisons with identical training observations; final four outcome observations retained.",
    10: "Training-matched contrasts; final four outcome observations excluded; resampling sensitivity.",
    11: "One-week seasonal-form comparisons; final four outcome observations excluded.",
    12: "One-week seasonal-form comparisons; final four outcome observations retained.",
}
WIDTHS = {
    1: [19, 26, 28, 28, 30, 22, 25],
    2: [23, 43, 86],
    3: [19, 15, 45, 10, 17, 20, 38],
    4: [19, 15, 36, 48, 10, 39],
    5: [19, 56, 19, 30, 30, 30, 30, 33],
    6: [19, 38, 49, 10, 39],
    7: [19, 48, 10, 18, 18, 25, 25],
    8: [19, 15, 37, 48, 10, 20, 20, 39],
    9: [19, 15, 37, 48, 10, 20, 20, 39],
    10: [19, 58, 15, 30, 30, 30, 30, 33],
    11: [19, 28, 27, 65, 10, 39],
    12: [19, 28, 27, 65, 10, 39],
}


def read_tsv(path):
    with path.open(newline="", encoding="utf-8") as f:
        return list(csv.reader(f, delimiter="\t"))


def value(text):
    normalized = text.replace("−", "-")
    if re.fullmatch(r"-?\d+", normalized):
        return int(normalized)
    if re.fullmatch(r"-?\d+\.\d+", normalized):
        return float(normalized)
    # Confidence intervals, date spans and annotated fractions are composite
    # scholarly labels; keep their complete published notation in one cell.
    return text


def line_count(text, width):
    return max(1, len(textwrap.wrap(str(text), max(8, int(width - 3)),
                                   break_long_words=False, break_on_hyphens=False)))


def merged_text(ws, row, end_col, text, *, bold=False, size=11, color=INK):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=end_col)
    cell = ws.cell(row, 1, text)
    cell.font = Font(name=FONT, size=size, bold=bold, color=color)
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    width = sum(ws.column_dimensions[get_column_letter(c)].width for c in range(1, end_col + 1))
    ws.row_dimensions[row].height = max(23, line_count(text, width * 0.88) * (size + 5) + 7)


def setup(ws):
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = 0.3
    ws.page_margins.right = 0.3
    ws.page_margins.top = 0.4
    ws.page_margins.bottom = 0.4
    ws.oddFooter.center.text = "Page &P of &N"
    ws.oddFooter.center.size = 9
    ws.oddFooter.center.font = FONT


def add_table(wb, sheet_name, title, rows, notes, widths, source, scope="", tab_color=None):
    ws = wb.create_sheet(sheet_name)
    setup(ws)
    if sheet_name == "Table S4":
        ws.page_setup.orientation = "portrait"
    n = len(rows[0])
    for j, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = width
    ws.cell(1, 1, "Contents").hyperlink = "#'Contents'!A1"
    ws.cell(1, 1).font = Font(name=FONT, size=10, color="245B88", underline="single")
    merged_text(ws, 2, n, title, bold=True, size=14)
    if scope:
        merged_text(ws, 3, n, scope, size=10, color=MUTED)
    merged_text(ws, 4, n, source, size=10, color=MUTED)
    ws.row_dimensions[5].height = 9
    first = 6
    last = first + len(rows) - 1
    for i, row in enumerate(rows, first):
        heights = []
        for j, text in enumerate(row, 1):
            cell = ws.cell(i, j, text if i == first else value(text))
            cell.font = Font(name=FONT, size=11, color=INK)
            cell.alignment = Alignment(horizontal="right" if isinstance(cell.value, (int, float))
                                       else "left", vertical="center", wrap_text=True, indent=1)
            if isinstance(cell.value, int):
                cell.number_format = "#,##0"
            elif isinstance(cell.value, float):
                cell.number_format = '0.000;−0.000;0.000'
            cell.border = Border(bottom=LINE)
            if i == first:
                cell.font = Font(name=FONT, size=11, bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor=NAVY)
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            elif (i - first) % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=GRAY)
            heights.append(line_count(text, widths[j - 1] - (3 if i != first else 0)) * 15 + 8)
        ws.row_dimensions[i].height = max(30 if i == first else 25, max(heights))
    table = Table(displayName=re.sub(r"\W", "", sheet_name) + "Data",
                  ref=f"A{first}:{get_column_letter(n)}{last}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False,
                                         showLastColumn=False, showRowStripes=True)
    ws.add_table(table)
    ws.freeze_panes = ("C7" if sheet_name in ["Table S5", "Table S10"] else
                       "D7" if n >= 6 else "B7")
    notes_start = last + 2
    merged_text(ws, notes_start, n, "Notes", bold=True)
    for i, note in enumerate(notes, notes_start + 1):
        merged_text(ws, i, n, note)
    ws.print_title_rows = "2:6"
    ws.print_area = f"A1:{get_column_letter(n)}{notes_start + len(notes)}"
    if tab_color:
        ws.sheet_properties.tabColor = tab_color
    return {"sheet":sheet_name, "header_row":first, "first_data_row":first + 1,
            "data_rows":len(rows) - 1, "columns":n, "notes_start":notes_start,
            "footnotes":notes, "headers":rows[0]}


def load_table_definitions():
    definitions=json.loads((ROOT/'data/table_definitions.json').read_text())
    tables=[]
    for n in range(1,13):
        item=definitions['tables'][str(n)]
        tables.append((n,item['title'],read_tsv(ROOT/'tables'/f'TableS{n}.tsv'),item['notes']))
    return definitions,tables


def main():
    definitions, tables = load_table_definitions()
    (REV / "tables").mkdir(exist_ok=True)
    wb = Workbook()
    cover = wb.active
    cover.title = "Contents"
    setup(cover)
    cover.page_setup.orientation = "portrait"
    cover.sheet_properties.tabColor = NAVY
    cover.column_dimensions["A"].width = 22
    cover.column_dimensions["B"].width = 140
    merged_text(cover, 2, 2, "Supplementary tables", bold=True, size=14)
    project_title = definitions["title"]
    merged_text(cover, 3, 2, project_title)
    merged_text(cover, 4, 2, "Hospitalization outcomes concern adults aged ≥65 years. The last observed outcome week ended July 4, 2026.", color=MUTED)
    merged_text(cover, 5, 2, "Select a table name below to open its worksheet. Table notes define the evaluation sample, training history and uncertainty.", color=MUTED)
    cover.cell(7, 1, "Table")
    cover.cell(7, 2, "Title")
    for cell in cover[7]:
        if cell.column <= 2:
            cell.font = Font(name=FONT, size=11, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=NAVY)
            cell.alignment = Alignment(vertical="center")
    metadata = []
    for number, title, rows, notes in tables:
        name = f"Table S{number}"
        if number == 1:
            notes = notes + ["wILI was used as a predictor only for influenza. Its RSV row describes source coverage."]
        metadata.append(add_table(wb, name, f"{name}. {title}", rows, notes, WIDTHS[number],
                                  f"Source: Supplementary materials, {name}.", SCOPES[number],
                                  "607D8B" if number >= 8 else None))
        row = number + 7
        cover.cell(row, 1, name).hyperlink = f"#'{name}'!A1"
        cover.cell(row, 1).font = Font(name=FONT, size=11, color="245B88", underline="single")
        cover.cell(row, 2, title).font = Font(name=FONT, size=11, color=INK)
        cover.cell(row, 2).alignment = Alignment(vertical="center", wrap_text=True)
        cover.row_dimensions[row].height = 30
    publication_rows = read_tsv(REV / "tables/Publication_timing.tsv")
    metadata.append(add_table(wb, "Publication timing", "Publication timing of surveillance sources",
                              publication_rows, ["Dates refer to the underlying surveillance week. Historical releases were not reconstructed."],
                              [28, 28, 73, 66], "Source: Supplementary Methods, End-of-series exclusion and publication timing.",
                              "Publication timing and use of retrieved surveillance snapshots."))
    cover.cell(20, 1, "Publication timing").hyperlink = "#'Publication timing'!A1"
    cover.cell(20, 1).font = Font(name=FONT, size=11, color="245B88", underline="single")
    cover.cell(20, 2, "Publication timing of surveillance sources").font = Font(name=FONT, size=11, color=INK)
    cover.row_dimensions[20].height = 30
    merged_text(cover, 22, 2, "Positive improvement favors the candidate model. A 95% confidence interval including zero does not establish improvement.", color=MUTED)
    merged_text(cover, 23, 2, "MAE, RMSE, mean bias, calibration intercepts and improvements are expressed in hospitalizations per 100,000 adults aged ≥65 years per week. Calibration slopes are dimensionless.", color=MUTED)
    merged_text(cover, 25, 2, "Abbreviations", bold=True)
    glossary = definitions["abbreviations"]
    glossary_items = glossary.rstrip(".").split("; ")
    if not any(item.startswith('RMSE,') for item in glossary_items):
        glossary_items.insert(9, "RMSE, root mean squared error")
    for row, item in enumerate(glossary_items, 26):
        term, definition = item.split(", ", 1)
        cover.cell(row, 1, term).font = Font(name=FONT, size=11, bold=True, color=INK)
        cover.cell(row, 2, definition).font = Font(name=FONT, size=11, color=INK)
        cover.cell(row, 2).alignment = Alignment(vertical="center", wrap_text=True)
        cover.row_dimensions[row].height = 25
    cover.print_area = f"A1:B{row}"
    wb.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True, calcMode="auto")
    wb.properties.title = "Supplementary tables for older-adult respiratory hospitalization forecasting"
    wb.properties.subject = project_title
    wb.properties.creator = ""
    wb.properties.lastModifiedBy = ""
    output = REV / "tables" / NAME
    wb.save(output)
    saved = load_workbook(output, data_only=False)
    expected_names = ["Contents"] + [f"Table S{n}" for n in range(1, 13)] + ["Publication timing"]
    assert saved.sheetnames == expected_names
    for item in metadata:
        ws = saved[item["sheet"]]
        source_name = "Publication_timing.tsv" if item["sheet"] == "Publication timing" else item["sheet"].replace(" ", "") + ".tsv"
        rows = read_tsv(REV / "tables" / source_name)
        for i, row in enumerate(rows, item["header_row"]):
            for j, original in enumerate(row, 1):
                expected = original if i == item["header_row"] else value(original)
                actual = ws.cell(i, j).value
                assert actual == expected, (item["sheet"], i, j, original, actual)
        assert len(ws.tables) == 1 and ws.freeze_panes
        for k, note in enumerate(item["footnotes"], item["notes_start"] + 1):
            assert ws.cell(k, 1).value == note
    assert not any(cell.data_type in ["f", "e"] for ws in saved for row in ws for cell in row)
    print(f'Created {output.name}: {len(saved.sheetnames)} worksheets.')


if __name__ == "__main__":
    main()
